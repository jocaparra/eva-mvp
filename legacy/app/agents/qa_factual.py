"""Validação factual de números e citações — parametrizada por artifact_type."""

from __future__ import annotations

from typing import List

from app.artifact_types import EXPECTED_HEADING_COUNT, EXPECTED_SLIDE_COUNT, REQUIRED_SECTIONS
from app.citations.matching import (
    AUDITABLE_FIELDS,
    CitationStatus,
    audit_financial_structured,
    audits_to_serializable,
)
from app.state import JobState


def resolve_audit_mode(state: JobState) -> str:
    """
    Política de auditoria:
    - full: deal com data room → QA factual obrigatório
    - legacy: POST /jobs sem deal_id → isento de QA factual (não auditável)
    """
    if state.get("audit_mode"):
        return str(state["audit_mode"])
    return "full" if state.get("deal_id") else "legacy"


def _artifact_type(state: JobState) -> str:
    return state.get("artifact_type") or "cim_pptx"


def validate_sections(state: JobState) -> List[str]:
    """Checagem estrutural por artifact_type (slides PPT ou headings DOCX)."""
    artifact_type = _artifact_type(state)
    required = REQUIRED_SECTIONS.get(artifact_type, [])
    issues: List[str] = []

    file_path = state.get("file_path") or state.get("ppt_path")
    if not file_path:
        return issues

    if artifact_type == "cim_pptx":
        expected = EXPECTED_SLIDE_COUNT.get(artifact_type, 1)
        try:
            from pptx import Presentation

            slide_count = len(Presentation(file_path).slides)
            if slide_count < expected:
                issues.append(
                    f"CIM PPT com {slide_count} slides; esperado ≥ {expected}."
                )
        except Exception:
            pass
        return issues

    if artifact_type == "memo_docx":
        expected_headings = EXPECTED_HEADING_COUNT.get(artifact_type, len(required))
        try:
            from docx import Document

            doc = Document(file_path)
            headings = [
                p.text.strip()
                for p in doc.paragraphs
                if p.style and p.style.name.startswith("Heading")
            ]
            if len(headings) < expected_headings:
                issues.append(
                    f"Memo DOCX com {len(headings)} headings; esperado ≥ {expected_headings} "
                    f"para {artifact_type}."
                )
            missing = [s for s in required if s != "Referências" and s not in headings]
            if missing:
                issues.append(
                    f"Memo sem seções obrigatórias: {', '.join(missing[:3])}."
                )
        except Exception:
            pass

    return issues


def validate_factual(state: JobState) -> tuple[List[str], dict]:
    """Valida citações de números-chave. Retorna (issues bloqueantes, field_audits)."""
    issues: List[str] = []
    financial = state.get("financial_structured") or {}
    citations = state.get("financial_citations") or []

    audits = audit_financial_structured(financial, citations)
    state_audits = audits_to_serializable(audits)

    for field in AUDITABLE_FIELDS:
        audit = audits[field]
        if audit.status == CitationStatus.UNCITED:
            issues.append(
                f"Campo '{field}' = '{audit.value}' sem citação no data room."
            )
        elif audit.status == CitationStatus.DIVERGENT:
            issues.append(
                f"Campo '{field}' = '{audit.value}' diverge dos números nas fontes."
            )

    return issues, state_audits


def run_factual_qa(state: JobState) -> dict:
    """Executa QA factual completo respeitando política legado vs deal."""
    issues: List[str] = []
    audit_mode = resolve_audit_mode(state)
    field_audits: dict = {}

    if state.get("error"):
        issues.append(state["error"])

    issues.extend(validate_sections(state))

    if audit_mode == "legacy":
        legacy_note = (
            "Auditoria factual não aplicável (fluxo legado sem deal_id — output não auditável)."
        )
        structural = [i for i in issues if i != legacy_note]
        return {
            "qa_passed": len(structural) == 0,
            "qa_issues": issues + [legacy_note],
            "audit_mode": audit_mode,
            "field_audits": {},
        }

    factual_issues, field_audits = validate_factual(state)
    issues.extend(factual_issues)

    return {
        "qa_passed": len(issues) == 0,
        "qa_issues": issues,
        "audit_mode": audit_mode,
        "field_audits": field_audits,
    }
